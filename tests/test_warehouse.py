"""Business invariants and concurrent clients sharing durable SQLite state."""
import asyncio
import io
from concurrent.futures import ThreadPoolExecutor

import httpx
import openpyxl
import pytest
from fastapi.testclient import TestClient

import main
from mockclient import MockTrueApiClient
from service import Service
from tokens import FileTokenStore
from trueapi import PRODUCT_INFO_PATH, TrueApiClient
from warehouse_excel import (canonical_full, decode_excel_string, export_shipment,
                             import_xlsx, make_xlsx, read_excel_strings, scan_data, template)
from warehouse_service import check_document, resolve_pns
from warehouse_store import WarehouseError, WarehouseStore

GTIN = '04640638345218'
ORG = '0000000001'
ORG2 = '0000000002'


def code(serial='INTRODUCEDQ100', tail='91EE12\x1d92ABCDcrypto+/='):
    return '01' + GTIN + '21' + serial + '\x1d' + tail


def workbook(kind='shipment', rows=None, org=ORG):
    metadata = [('Поле', 'Значение'), ('Номер документа', 'Тест/001'), ('Дата документа', '2026-10-21'), ('ИНН организации', org)]
    if kind == 'shipment':
        metadata += [('Получатель', 'Получатель'), ('ИНН получателя', '1234567890')]
    cols = ['PN', 'GTIN', 'Количество'] if kind == 'shipment' else ['PN', 'GTIN', 'Код маркировки', 'Количество в КМ']
    return make_xlsx({'Документ': metadata, 'Товары': [cols] + (rows or [['PN-1', GTIN, 100]])})


@pytest.fixture
def store(tmp_path):
    return WarehouseStore(tmp_path / 'warehouse.sqlite3')


@pytest.fixture
def svc(tmp_path):
    tokens = FileTokenStore(tmp_path / 'tokens.json')
    tokens.set(ORG, 'token1')
    tokens.set(ORG2, 'token2')
    return Service(MockTrueApiClient(), tokens)


def document(store, kind='shipment', quantity=100, rows=None, start=True):
    if kind == 'receipt' and rows is None:
        rows = [['PN-1', GTIN, code(), quantity]]
    data = import_xlsx(workbook(kind, rows or [['PN-1', GTIN, quantity]]), kind)
    d = store.create(*data)
    return store.start(d['id']) if start else d


def add(store, d, raw=None):
    return store.add_scan(d['id'], scan_data(raw or code()))


@pytest.mark.parametrize('kind', ['shipment', 'receipt'])
def test_shared_scans_and_restart(store, svc, kind):
    d = document(store, kind)
    add(store, d)
    other = WarehouseStore(store.path)
    assert len(other.get(d['id'])['scans']) == 1
    assert other.get(d['id'])['result'] is None
    result = asyncio.run(check_document(other, svc, d['id']))
    assert result['result']['ok'] and result['can_complete']
    other.complete(d['id'])
    restarted = WarehouseStore(store.path)
    assert restarted.get(d['id'])['status'] == 'completed'
    with pytest.raises(WarehouseError):
        add(restarted, d, code('INTRODUCEDQ101'))
    with pytest.raises(WarehouseError):
        restarted.delete_scan(d['id'], result['scans'][0]['id'])


def test_atomic_duplicate_across_devices(store):
    d = document(store)
    a, b = WarehouseStore(store.path), WarehouseStore(store.path)
    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(lambda s: add(s, d), [a, b]))
    assert sorted(r['duplicate'] for r in results) == [False, True]
    assert len(store.get(d['id'])['scans']) == 1
    assert store.get(d['id'])['version'] == d['version'] + 1


@pytest.mark.asyncio
async def test_russian_scan_matches_receipt_and_is_stored_as_english(store, svc):
    english = ']d2010464063834521821AbC"d/ef\x1d91EE12\x1d92+Ab/Cd=='
    russian = 'ъв2010464063834521821ФиСЭв.уа\x1d91УУ12\x1d92+Фи.Св=='
    d = document(store, 'receipt', rows=[['PN', GTIN, english, 100]])
    first = add(store, d, russian)
    assert first['scan']['raw_code'] == english
    assert add(store, d, english)['duplicate']
    checked = await check_document(store, svc, d['id'])
    assert checked['can_complete'] and checked['result']['ok']


@pytest.mark.parametrize('kind', ['receipt', 'shipment'])
@pytest.mark.asyncio
async def test_completed_documents_hidden_from_warehouse_and_reopened_are_visible(store, svc, kind):
    d = document(store, kind)
    add(store, d)
    await check_document(store, svc, d['id'])
    assert any(row['id'] == d['id'] for row in store.list())
    store.complete(d['id'])
    assert not store.list()
    assert not store.list(status='completed')
    assert [row['id'] for row in store.list(admin=True, status='completed')] == [d['id']]
    store.start(d['id'], reopen=True)
    assert [row['id'] for row in store.list()] == [d['id']]


def test_concurrent_distinct_scans_are_not_lost(store):
    d = document(store)
    with ThreadPoolExecutor(8) as pool:
        list(pool.map(lambda i: add(WarehouseStore(store.path), d, code(f'INTRODUCEDQ{i}')), range(1, 21)))
    assert len(store.get(d['id'])['scans']) == 20


@pytest.mark.parametrize('kind', ['shipment', 'receipt'])
def test_scan_is_local_only_and_does_not_recalculate(store, svc, kind):
    d = document(store, kind)
    async def forbidden(*args, **kwargs):
        pytest.fail('Network called during scanning or receipt reconciliation')
    svc.client.info = forbidden
    svc.client.product_info = forbidden
    add(store, d)
    assert store.get(d['id'])['result'] is None
    if kind == 'receipt':
        checked = asyncio.run(check_document(store, svc, d['id']))
        assert checked['can_complete']
        before = checked['result']
        add(store, d, code('OTHERQ100'))
        after = store.get(d['id'])
        assert after['result'] == before and not after['check_current']


@pytest.mark.parametrize('status,ok', [('INTRODUCED', True), ('APPLIED', False), ('EMITTED', False),
                                      ('RETIRED', False), ('UNKNOWN', False), (None, False)])
def test_only_introduced_allowed(store, svc, status, ok):
    d = document(store)
    add(store, d)
    original = svc.client.info
    async def info(cises, token):
        records = await original(cises, token)
        for r in records:
            r['cisInfo']['status'] = status
        return records
    svc.client.info = info
    result = asyncio.run(check_document(store, svc, d['id']))
    assert result['can_complete'] is ok
    if not ok:
        with pytest.raises(WarehouseError):
            store.complete(d['id'])


@pytest.mark.parametrize('owner,strict,ok', [(ORG, True, True), (ORG2, True, False),
                                          (ORG2, False, True), ('7777777777', False, False),
                                          (None, False, False)])
def test_owner_policy(store, svc, owner, strict, ok):
    d = document(store)
    d = store.edit(d['id'], d['version'], strict_owner=strict)
    add(store, d)
    original = svc.client.info
    async def info(cises, token):
        records = await original(cises, token)
        for r in records:
            r['cisInfo']['ownerInn'] = owner
        return records
    svc.client.info = info
    assert asyncio.run(check_document(store, svc, d['id']))['can_complete'] is ok


def test_alternative_token_can_resolve_owner(store, svc):
    d = document(store)
    add(store, d)
    original = svc.client.info
    async def info(cises, token):
        records = await original(cises, token)
        if token == 'token1':
            for r in records:
                r['cisInfo']['ownerInn'] = None
        return records
    svc.client.info = info
    assert asyncio.run(check_document(store, svc, d['id']))['can_complete']


@pytest.mark.parametrize('value', [None, 0, -1, 1.5, 'wrong', True])
def test_unknown_or_invalid_quantity_blocks(store, svc, value):
    d = document(store)
    add(store, d)
    original = svc.client.info
    async def info(cises, token):
        records = await original(cises, token)
        for r in records:
            r['cisInfo']['quantityInPack'] = value
        return records
    svc.client.info = info
    checked = asyncio.run(check_document(store, svc, d['id']))
    assert not checked['can_complete'] and checked['result']['rows'][0]['actual'] == 0


@pytest.mark.parametrize('quantities,expected,ok', [([100,200],300,True), ([100],300,False), ([100,200],100,False)])
def test_sum_quantity_not_number_of_codes(store, svc, quantities, expected, ok):
    d = document(store, quantity=expected)
    for i,q in enumerate(quantities):
        add(store, d, code(f'INTRODUCEDQ{q}S{i}'))
    result = asyncio.run(check_document(store, svc, d['id']))
    assert result['result']['rows'][0]['actual'] == sum(quantities)
    assert result['can_complete'] is ok


def test_extra_gtin_blocks_even_when_planned_quantity_matches(store, svc):
    d = document(store)
    add(store, d)
    add(store, d, code('INTRODUCEDQ100').replace(GTIN, '04646063126927'))
    result = asyncio.run(check_document(store, svc, d['id']))
    assert not result['can_complete']


def test_api_failure_blocks_completion(store, svc):
    d = document(store)
    add(store, d)
    async def fail(*args):
        raise httpx.ConnectError('unavailable')
    svc.client.info = fail
    result = asyncio.run(check_document(store, svc, d['id']))
    assert not result['can_complete']
    assert result['result']['codes'][0]['errors']


def test_api_gtin_mismatch_blocks(store, svc):
    d = document(store)
    add(store, d)
    original = svc.client.info
    async def info(cises, token):
        records = await original(cises, token)
        records[0]['cisInfo']['gtin'] = '00000000000000'
        return records
    svc.client.info = info
    assert not asyncio.run(check_document(store, svc, d['id']))['can_complete']


def test_scan_during_network_check_makes_result_stale(store, svc):
    async def scenario():
        d = document(store)
        add(store, d)
        entered, release = asyncio.Event(), asyncio.Event()
        original = svc.client.info
        async def slow_info(cises, token):
            entered.set()
            await release.wait()
            return await original(cises, token)
        svc.client.info = slow_info
        task = asyncio.create_task(check_document(store, svc, d['id']))
        await entered.wait()
        assert store.get(d['id'])['checking']
        with pytest.raises(WarehouseError):
            store.begin_check(d['id'])
        add(WarehouseStore(store.path), d, code('INTRODUCEDQ200'))
        release.set()
        result = await task
        assert result['result']['ok']  # the original snapshot matches
        assert not result['check_current'] and not result['can_complete']
        assert len(result['scans']) == 2
        with pytest.raises(WarehouseError):
            store.complete(d['id'])
    asyncio.run(scenario())


def test_delete_and_owner_mode_invalidate_check(store, svc):
    d = document(store)
    added = add(store, d)
    checked = asyncio.run(check_document(store, svc, d['id']))
    assert checked['can_complete']
    updated = store.edit(d['id'], checked['version'], strict_owner=False)
    assert not updated['check_current']
    checked = asyncio.run(check_document(store, svc, d['id']))
    deleted = store.delete_scan(d['id'], added['scan']['id'])
    assert not deleted['can_complete'] and deleted['result'] == checked['result']


def test_complete_requires_explicit_check_and_freezes_all_devices(store):
    d = document(store)
    add(store, d)
    with pytest.raises(WarehouseError):
        store.complete(d['id'])


def test_receipt_requires_exact_set_and_full_crypto(store, svc):
    d = document(store, 'receipt', rows=[['PN',GTIN,code('A'),100],['PN',GTIN,code('B'),100]])
    add(store, d, code('A'))
    result = asyncio.run(check_document(store, svc, d['id']))
    assert not result['can_complete'] and len(result['result']['missing']) == 1
    add(store, d, code('B', tail='91EE12\x1d92DIFFERENT'))
    result = asyncio.run(check_document(store, svc, d['id']))
    assert not result['can_complete']
    assert result['result']['codes'][1]['errors']


def test_receipt_unknown_code_never_counts(store, svc):
    d = document(store, 'receipt')
    add(store, d, code('INTRODUCEDQ200'))
    result = asyncio.run(check_document(store, svc, d['id']))
    assert not result['can_complete'] and result['result']['rows'][0]['actual'] == 0


def test_receipt_prefix_and_gs_aliases_match(store, svc):
    d = document(store, 'receipt')
    add(store, d, ']d2' + code().replace('\x1d','{GS}'))
    assert asyncio.run(check_document(store, svc, d['id']))['can_complete']


def test_excel_preserves_full_dm_gs_and_literals(store, svc):
    raw = code(tail='91EE12\x1d92+/=_x001D_&<>')
    d = document(store)
    add(store, d, raw)
    asyncio.run(check_document(store, svc, d['id']))
    completed = store.complete(d['id'])
    content = export_shipment(completed)
    book = openpyxl.load_workbook(io.BytesIO(content))
    cell = book.active['C2']
    assert cell.data_type == 's'
    assert read_excel_strings(content)['Коды маркировки']['C2'] == raw
    assert book.active['B2'].value == GTIN
    received = import_xlsx(workbook('receipt', [['PN',GTIN,raw,100]]), 'receipt')
    assert received[2][0]['raw_code'] == raw


def test_pn_required_before_export_but_can_be_supplied_after_completion(store, svc):
    d = document(store, rows=[['', GTIN, 100]])
    async def no_pn(*args): return []
    svc.client.product_info = no_pn
    add(store, d)
    asyncio.run(check_document(store, svc, d['id']))
    completed = store.complete(d['id'])
    with pytest.raises(WarehouseError): export_shipment(completed)
    updated = store.edit(d['id'], completed['version'], pns={GTIN:'MANUAL-PN'})
    assert updated['check_current']
    assert export_shipment(updated)


def test_resolve_only_documented_model_and_never_product_name(svc):
    assert asyncio.run(resolve_pns(svc,[GTIN],ORG))[0][GTIN].startswith('MOCK-PN-')
    async def no_model(*args): return [{'gtin':GTIN,'name':'NOT-PN','productName':'NOT-PN'}]
    svc.client.product_info = no_model
    found, warnings = asyncio.run(resolve_pns(svc,[GTIN],ORG))
    assert found == {} and warnings
    async def failure(*args): raise httpx.ConnectError('offline')
    svc.client.product_info = failure
    assert asyncio.run(resolve_pns(svc,[GTIN],ORG))[0] == {}


@pytest.mark.parametrize('rows,contains', [
    ([['PN',GTIN,0]],'Количество'),([['PN',GTIN,1.2]],'Количество'),
    ([['PN',int(GTIN),100]],'Текст'),([['PN','04640638345219',100]],'контрольной'),
    ([['PN',GTIN,100],['OTHER',GTIN,100]],'разные PN'),
])
def test_import_row_errors(rows, contains):
    with pytest.raises(WarehouseError, match=contains):
        import_xlsx(workbook(rows=rows),'shipment')


def test_repeated_gtin_rows_merged():
    _, lines, _ = import_xlsx(workbook(rows=[['',GTIN,100],['PN',GTIN,200]]),'shipment')
    assert lines == [dict(gtin=GTIN,pn='PN',pn_source='excel',quantity=300)]


def test_receipt_duplicate_expected_codes_rejected():
    with pytest.raises(WarehouseError,match='строка 3.*Дубликат'):
        import_xlsx(workbook('receipt',[['PN',GTIN,code(),100],['PN',GTIN,']d2'+code(),100]]),'receipt')


def test_receipt_gtin_mismatch_and_no_sender_required():
    header, _, _ = import_xlsx(workbook('receipt',[['PN',GTIN,code(),100]]),'receipt')
    assert not header['recipient_inn']
    with pytest.raises(WarehouseError,match='не соответствует'):
        import_xlsx(workbook('receipt',[['PN','04646063126927',code(),100]]),'receipt')


def test_template_format_and_invalid_file():
    for kind in ('shipment','receipt'):
        book = openpyxl.load_workbook(io.BytesIO(template(kind)))
        assert book.sheetnames == ['Документ','Товары']
        assert book['Товары'].column_dimensions['B'].number_format == '@'
    with pytest.raises(WarehouseError): import_xlsx(b'not xlsx','shipment')


def test_formula_cells_rejected_and_export_not_formula():
    book = openpyxl.load_workbook(io.BytesIO(workbook()))
    book['Товары']['A2'] = '=1+1'
    stream = io.BytesIO(); book.save(stream)
    with pytest.raises(WarehouseError,match='формулы'):
        import_xlsx(stream.getvalue(),'shipment')
    out = openpyxl.load_workbook(io.BytesIO(make_xlsx({'X':[['PN'],['=DANGEROUS()']]})))
    assert out.active['A2'].data_type == 's'


def test_admin_edit_optimistic_version_and_reopen(store, svc):
    d = document(store)
    add(store, d)
    with pytest.raises(WarehouseError,match='изменён'):
        store.edit(d['id'], d['version'], strict_owner=False)
    checked = asyncio.run(check_document(store, svc, d['id']))
    store.complete(d['id'])
    reopened = store.start(d['id'],reopen=True)
    assert not reopened['can_complete'] and not reopened['check_current']
    assert len(reopened['scans']) == 1


@pytest.fixture
def clients(monkeypatch, store, svc):
    monkeypatch.setattr(main,'warehouse',store)
    monkeypatch.setattr(main,'svc',svc)
    monkeypatch.setattr(main.settings,'admin_key','test-admin')
    return TestClient(main.app), TestClient(main.app)


def test_admin_endpoints_require_header_not_url(clients):
    a,_ = clients
    base = '/api/admin/warehouse'
    for path in ['/documents','/settings','/templates/shipment']:
        assert a.get(base+path).status_code == 403
        assert a.get(base+path+'?key=test-admin').status_code == 403
        assert a.get(base+path,headers={'X-Admin-Key':'test-admin'}).status_code == 200
    assert a.post(base+'/documents/a/start').status_code == 403
    assert a.patch(base+'/documents/a',json={'version':0}).status_code == 403
    assert a.get('/admin/warehouse').status_code == 200
    assert 'test-admin' not in a.get('/admin/warehouse').text


def test_http_import_pn_and_shared_workflow(clients):
    a,b = clients
    headers = {'X-Admin-Key':'test-admin'}
    created = a.post('/api/admin/warehouse/documents', headers=headers,
                     data={'kind':'shipment'},files={'file':('shipment.xlsx',workbook(rows=[['',GTIN,100]]))})
    assert created.status_code == 200, created.text
    d = created.json(); path = '/api/warehouse/documents/' + d['id']
    assert d['lines'][0]['pn_source'] == 'trueapi'
    assert b.get(path).status_code == 404
    assert a.post('/api/admin/warehouse/documents/'+d['id']+'/start',headers=headers).status_code == 200
    assert a.post(path+'/scans',json={'code':code()}).json()['duplicate'] is False
    assert len(b.get(path).json()['scans']) == 1
    assert b.post(path+'/scans',json={'code':']d2'+code()}).json()['duplicate'] is True
    assert b.post(path+'/complete').status_code == 409
    assert b.post(path+'/check').json()['can_complete']
    assert a.post(path+'/complete').json()['status'] == 'completed'
    assert b.post(path+'/scans',json={'code':code('INTRODUCEDQ200')}).status_code == 409
    assert a.get('/api/admin/warehouse/documents/'+d['id']+'/export',headers=headers).status_code == 200


def test_product_info_uses_documented_endpoint_and_chunks():
    calls=[]
    def handler(req):
        import json
        body=json.loads(req.content)
        calls.append((req.url.path,body,req.headers['Authorization']))
        return httpx.Response(200,json={'results':[{'gtin':g,'model':'PN'} for g in body['gtins']]})
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            client=TrueApiClient('https://example.invalid',http=http)
            result=await client.product_info([GTIN]*1001,'secret')
            assert len(result)==1001
    asyncio.run(scenario())
    assert [len(c[1]['gtins']) for c in calls]==[1000,1]
    assert calls[0][0]==PRODUCT_INFO_PATH
    assert calls[0][2]=='Bearer secret'


def test_plan_edits_invalidate_check_and_metadata_has_optimistic_revision(store, svc):
    d = document(store)
    add(store,d)
    checked = asyncio.run(check_document(store,svc,d['id']))
    assert checked['can_complete']
    changed = store.edit(d['id'],checked['version'],quantities={GTIN:200},revision=checked['revision'])
    assert not changed['check_current'] and not changed['can_complete']
    with pytest.raises(WarehouseError):
        store.edit(d['id'],changed['version'],quantities={GTIN:0})
    first_edit = store.edit(d['id'],changed['version'],pns={GTIN:'NEW'},revision=changed['revision'])
    with pytest.raises(WarehouseError,match='изменён'):
        store.edit(d['id'],changed['version'],pns={GTIN:'STALE'},revision=changed['revision'])
    assert store.get(d['id'])['lines'][0]['pn'] == 'NEW'


def test_admin_key_unicode_is_rejected_and_token_page_contains_no_key(clients):
    a,_=clients
    assert a.get('/api/admin/warehouse/documents', headers={'X-Admin-Key': 'wrong'}).status_code==403
    assert main._check_admin('неверный ключ') is False
    page=a.get('/admin/token?key=test-admin')
    assert page.status_code==200 and 'test-admin' not in page.text
    assert a.post('/admin/token?key=test-admin', data={'inn':ORG,'uuid':'uuid','signature':'signature'}).status_code==403
    assert a.post('/admin/token',headers={'X-Admin-Key':'test-admin'},data={'inn':ORG,'uuid':'uuid','signature':'signature'}).json()['ok']


def test_admin_session_survives_refresh_and_covers_all_sections(clients):
    a,b = clients
    login = a.post('/api/admin/session',json={'key':'test-admin'})
    assert login.status_code == 200
    cookie = login.headers['set-cookie'].lower()
    assert 'httponly' in cookie and 'samesite=strict' in cookie and 'max-age=28800' in cookie
    assert 'test-admin' not in cookie
    assert a.get('/api/admin/session').status_code == 200
    assert a.get('/api/admin/warehouse/documents').status_code == 200
    assert b.get('/api/admin/warehouse/documents').status_code == 403
    assert a.get('/admin').status_code == 200
    assert a.get('/admin/token').url.path == '/admin'
    assert a.get('/admin/token').url.query == b'section=token'
    r = a.post('/admin/token', headers={'X-Admin-Request':'1'},
               data={'inn':ORG,'uuid':'uuid','signature':'signature'})
    assert r.json()['ok']
    assert a.delete('/api/admin/session',headers={'X-Admin-Request':'1'}).status_code == 200
    assert a.get('/api/admin/warehouse/documents').status_code == 403


def test_cookie_session_requires_csrf_header_for_writes(clients):
    a,_ = clients
    a.post('/api/admin/session',json={'key':'test-admin'})
    for path,method,body in [('/api/admin/warehouse/documents/missing/start','POST',None),
                             ('/api/admin/warehouse/documents/missing','DELETE',{'version':0,'revision':0})]:
        assert a.request(method,path,json=body).status_code == 403
    assert a.post('/admin/token',data={'inn':ORG,'uuid':'uuid','signature':'signature'}).status_code == 403
    assert a.delete('/api/admin/session').status_code == 403
    assert a.get('/api/admin/session').status_code == 200


def test_session_expiration_tampering_and_key_rotation(clients, monkeypatch):
    import admin_auth
    a,_ = clients
    a.post('/api/admin/session',json={'key':'test-admin'})
    value = a.cookies.get(admin_auth.COOKIE_NAME)
    assert admin_auth.valid_session(value,'test-admin')
    assert not admin_auth.valid_session(value+'x','test-admin')
    assert not admin_auth.valid_session('invalid','test-admin')
    assert not admin_auth.valid_session(value,'change-me-in-production')
    monkeypatch.setattr(main.settings,'admin_key','rotated-key')
    assert a.get('/api/admin/session').status_code == 403
    monkeypatch.setattr(main.settings,'admin_key','test-admin')
    expires = int(value.split('.')[0])
    monkeypatch.setattr(admin_auth.time,'time',lambda:expires + 1)
    assert a.get('/api/admin/session').status_code == 403


def test_session_fail_closed_and_https_cookie(clients,monkeypatch):
    a,_ = clients
    assert a.post('/api/admin/session',json={'key':'wrong'}).status_code == 403
    assert a.post('/api/admin/session',json={'key':'неверный ключ'}).status_code == 403
    https = TestClient(main.app,base_url='https://testserver')
    assert '; Secure' in https.post('/api/admin/session',json={'key':'test-admin'}).headers['set-cookie']
    monkeypatch.setattr(main.settings,'admin_key','change-me-in-production')
    assert a.post('/api/admin/session',json={'key':'change-me-in-production'}).status_code == 403


@pytest.mark.parametrize('kind',['shipment','receipt'])
@pytest.mark.parametrize('status',['draft','active','completed'])
def test_admin_document_deletion_hides_and_blocks_operations_but_keeps_history(store,svc,kind,status):
    d = document(store,kind,start=status!='draft')
    if status != 'draft': add(store,d)
    if status == 'completed':
        asyncio.run(check_document(store,svc,d['id']))
        store.complete(d['id'])
    d = store.get(d['id'])
    store.delete(d['id'],d['version'],d['revision'])
    assert not store.list(admin=True) and not store.list()
    with pytest.raises(WarehouseError): store.get(d['id'])
    with pytest.raises(WarehouseError): add(store,d)
    with pytest.raises(WarehouseError): store.start(d['id'],reopen=True)
    with pytest.raises(WarehouseError): store.complete(d['id'])
    with pytest.raises(WarehouseError): store.edit(d['id'],d['version'],pns={GTIN:'PN'})
    with store.connection() as c:
        row=c.execute('SELECT deleted_at FROM warehouse_documents WHERE id=?',(d['id'],)).fetchone()
        assert row['deleted_at']
        if status != 'draft':
            assert c.execute('SELECT count(*) FROM warehouse_scans WHERE document_id=?',(d['id'],)).fetchone()[0] == 1
    assert not WarehouseStore(store.path).list(admin=True)


def test_deletion_rejects_stale_confirmation_and_running_check(store):
    d = document(store)
    add(store,d)
    with pytest.raises(WarehouseError,match='изменён'):
        store.delete(d['id'],d['version'],d['revision'])
    updated=store.get(d['id'])
    check_id,_=store.begin_check(d['id'])
    with pytest.raises(WarehouseError,match='Дождитесь'):
        store.delete(d['id'],updated['version'],updated['revision'])
    store.cancel_check(d['id'],check_id)
    store.delete(d['id'],updated['version'],updated['revision'])


def test_delete_api_requires_admin_and_stops_other_device(clients,store):
    a,b=clients
    d=document(store)
    path='/api/admin/warehouse/documents/'+d['id']
    body={'version':d['version'],'revision':d['revision']}
    assert b.request('DELETE',path,json=body).status_code==403
    a.post('/api/admin/session',json={'key':'test-admin'})
    assert a.request('DELETE',path,json=body).status_code==403
    result=a.request('DELETE',path,json=body,headers={'X-Admin-Request':'1'})
    assert result.status_code==200 and result.json()['deleted']
    public='/api/warehouse/documents/'+d['id']
    assert b.get(public).status_code==404
    assert b.post(public+'/scans',json={'code':code()}).status_code==404
    assert b.post(public+'/check').status_code==404
    assert b.post(public+'/complete').status_code==404
    assert a.get(path).json()['deleted_at']
    assert a.get(path+'/export').status_code==404


def test_document_filters_distinguish_scans_from_launch_status(store):
    draft=document(store,start=False)
    empty=document(store)
    scanned=document(store);add(store,scanned)
    assert {d['id'] for d in store.list(scan_state='unscanned')}=={empty['id']}
    assert {d['id'] for d in store.list(scan_state='scanned')}=={scanned['id']}
    assert {d['id'] for d in store.list(admin=True,status='draft')}=={draft['id']}
    assert not store.list(status='draft')
    assert len(store.list(org_inn=ORG))==2
    assert not store.list(org_inn=ORG2)


def test_admin_search_dates_and_trash_filters(store):
    a=document(store)
    b=document(store,start=False)
    with store.connection(True) as c:
        c.execute("UPDATE warehouse_documents SET number='Тест Альфа',recipient_name='Клиент Бета',date='2026-10-20' WHERE id=?",(a['id'],))
    assert [d['id'] for d in store.list(admin=True,query='КЛИЕНТ бета')]==[a['id']]
    assert [d['id'] for d in store.list(admin=True,query='альфа')]==[a['id']]
    assert len(store.list(admin=True,date_from='2026-10-21',date_to='2026-10-21'))==1
    assert not store.list(admin=True,query="%' OR 1=1 --")
    store.delete(b['id'],b['version'],b['revision'])
    assert [d['id'] for d in store.list(admin=True,trash=True)]==[b['id']]
    assert [d['id'] for d in store.list(trash=True)]==[a['id']]  # worker cannot access trash


@pytest.mark.parametrize('kind',['shipment','receipt'])
@pytest.mark.parametrize('status',['draft','active','completed'])
def test_restore_preserves_scans_and_original_status(store,svc,kind,status):
    d=document(store,kind,start=status!='draft')
    if status!='draft':add(store,d)
    if status=='completed':
        asyncio.run(check_document(store,svc,d['id']));store.complete(d['id'])
    before=store.get(d['id']);store.delete(d['id'],before['version'],before['revision'])
    trashed=store.get(d['id'],include_deleted=True)
    assert not trashed['can_complete']
    with pytest.raises(WarehouseError):store.restore(d['id'],before['version'],before['revision'])
    restored=store.restore(d['id'],trashed['version'],trashed['revision'])
    assert restored['status']==status and restored['scans']==before['scans']
    assert not restored['deleted_at'] and not restored['can_complete']
    assert not store.list(admin=True,trash=True)
    if status=='completed' and kind=='shipment':assert export_shipment(restored)
    with pytest.raises(WarehouseError):store.restore(d['id'],restored['version'],restored['revision'])


def failed_check(store,svc):
    d=document(store,quantity=200);add(store,d)
    return asyncio.run(check_document(store,svc,d['id']))


def test_problem_report_is_snapshot_with_operations_and_resolves_explicitly(store,svc):
    d=failed_check(store,svc)
    assert not d['result']['ok']
    issue=store.report_issue(d['id'],d['version'],d['result_at'],'Не хватает товара')
    assert store.get(d['id'])['reported_current_result']
    assert store.list(admin=True,issue_state='open')[0]['id']==d['id']
    assert not store.list(admin=True,issue_state='closed')
    store.delete_scan(d['id'],d['scans'][0]['id'])
    detail=store.issue(d['id'],issue['id'])
    assert detail['snapshot']['scans']==d['scans']
    assert detail['snapshot']['result']==d['result']
    events=store.issue_events(d['id'],issue['id'])['events']
    assert {'created','started','scan_added','check_started','checked','scan_removed','issue_reported'} <= {e['action'] for e in events}
    removed=next(e for e in events if e['action']=='scan_removed')
    assert removed['details']['raw_code']==d['scans'][0]['raw_code']
    with pytest.raises(WarehouseError):store.resolve_issue(d['id'],issue['id'],' ')
    store.resolve_issue(d['id'],issue['id'],'Склад заменил товар')
    assert not store.get(d['id'])['open_issues']
    assert store.list(admin=True,issue_state='closed')[0]['id']==d['id']
    assert store.issue(d['id'],issue['id'])['resolution']=='Склад заменил товар'


def test_report_deduplication_across_devices(store,svc):
    d=failed_check(store,svc)
    with ThreadPoolExecutor(2) as pool:
        reports=list(pool.map(lambda _:WarehouseStore(store.path).report_issue(d['id'],d['version'],d['result_at']),range(2)))
    assert sorted(r['duplicate'] for r in reports)==[False,True]
    assert reports[0]['id']==reports[1]['id']
    assert len(store.issues(d['id']))==1


def test_report_rejects_stale_unchecked_and_successful_results(store,svc):
    d=document(store)
    with pytest.raises(WarehouseError):store.report_issue(d['id'],d['version'],'none')
    add(store,d);ok=asyncio.run(check_document(store,svc,d['id']))
    with pytest.raises(WarehouseError):store.report_issue(d['id'],ok['version'],ok['result_at'])
    failed=failed_check(store,svc)
    add(store,failed,code('INTRODUCEDQ50'))
    with pytest.raises(WarehouseError):store.report_issue(failed['id'],failed['version'],failed['result_at'])
    fresh=asyncio.run(check_document(store,svc,failed['id']))
    with pytest.raises(WarehouseError):store.report_issue(failed['id'],fresh['version'],failed['result_at'])


def test_successful_recheck_does_not_clear_responsible_attention(store,svc):
    d=failed_check(store,svc)
    report=store.report_issue(d['id'],d['version'],d['result_at'])
    add(store,d,code('INTRODUCEDQ100B'))
    good=asyncio.run(check_document(store,svc,d['id']))
    assert good['can_complete'] and good['open_issues']==1
    assert not store.issue(d['id'],report['id'])['resolved_at']
    restarted=WarehouseStore(store.path)
    assert restarted.get(d['id'])['open_issues']==1


def test_problem_events_paginate_without_duplicates(store,svc):
    d=failed_check(store,svc)
    issue=store.report_issue(d['id'],d['version'],d['result_at'])
    first=store.issue_events(d['id'],issue['id'],limit=3)
    second=store.issue_events(d['id'],issue['id'],before=first['next_before'],limit=3)
    assert len(first['events'])==3
    assert not {e['id'] for e in first['events']} & {e['id'] for e in second['events']}


def test_filter_restore_and_problem_apis(clients,store,svc):
    a,b=clients;h={'X-Admin-Key':'test-admin'}
    d=failed_check(store,svc);path='/api/warehouse/documents/'+d['id'];admin='/api/admin/warehouse/documents/'+d['id']
    assert b.get('/api/warehouse/documents?scan_state=invalid').status_code==422
    assert a.get('/api/admin/warehouse/documents?date_from=2026-10-22&date_to=2026-10-21',headers=h).status_code==422
    report=b.post(path+'/issues',json={'version':d['version'],'result_at':d['result_at'],'message':'Нужна помощь'})
    assert report.status_code==200
    issue=report.json()['id']
    assert b.get(admin+'/issues').status_code==403
    assert b.get(admin+'/issues/'+issue).status_code==403
    assert a.get(admin+'/issues/'+issue,headers=h).json()['snapshot']['scans']==d['scans']
    assert a.get(admin+'/issues/'+issue+'/events',headers=h).json()['events']
    assert b.post(admin+'/issues/'+issue+'/resolve',json={'resolution':'Ок'}).status_code==403
    fresh=store.get(d['id']);store.delete(d['id'],fresh['version'],fresh['revision'])
    trashed=store.get(d['id'],include_deleted=True);body={'version':trashed['version'],'revision':trashed['revision']}
    assert b.post(admin+'/restore',json=body).status_code==403
    assert b.get('/api/warehouse/documents?trash=true').json()['documents']==[]
    assert a.get('/api/admin/warehouse/documents?trash=true&issue_state=open',headers=h).json()['documents'][0]['id']==d['id']
    assert a.post(admin+'/restore',headers=h,json=body).status_code==200
    assert a.post(admin+'/issues/'+issue+'/resolve',headers=h,json={'resolution':'Разобрано'}).status_code==200


def test_receipt_issue_keeps_missing_codes_without_any_true_api_call(store,svc):
    async def forbidden(*args,**kwargs):pytest.fail('Receipt/reporting must stay local')
    svc.client.info=forbidden;svc.client.product_info=forbidden
    d=document(store,'receipt')
    failed=asyncio.run(check_document(store,svc,d['id']))
    report=store.report_issue(d['id'],failed['version'],failed['result_at'],'Ожидаемый товар не приехал')
    detail=store.issue(d['id'],report['id'])
    assert detail['snapshot']['result']['missing'][0]['cis']==scan_data(code())['cis']
    assert detail['snapshot']['expected'][0]['raw_code']==code()
    assert not detail['snapshot']['scans']


def test_return_to_warehouse_preserves_check_and_does_not_block_work(store,svc):
    d=failed_check(store,svc)
    issue=store.report_issue(d['id'],d['version'],d['result_at'])
    add(store,d,code('INTRODUCEDQ100B'))
    checked=asyncio.run(check_document(store,svc,d['id']))
    assert checked['can_complete']
    reply=store.return_to_warehouse(d['id'],issue['id'],'Исправлено, продолжайте','request1')
    after=store.get(d['id'])
    assert after['version']==checked['version'] and after['check_current'] and after['can_complete']
    assert after['status']=='active' and after['open_issues']==0 and after['pending_returns']==1
    assert after['warehouse_replies'][0]['message']=='Исправлено, продолжайте'
    assert store.issue(d['id'],issue['id'])['resolved_at']
    assert store.list()[0]['pending_returns']==1
    # Completing is permitted before anyone acknowledges the reply.
    assert store.complete(d['id'])['status']=='completed'
    with pytest.raises(WarehouseError):store.return_to_warehouse(d['id'],issue['id'],'Ещё один ответ','completed')
    assert store.acknowledge_return(d['id'],reply['id'])['acknowledged']
    assert not store.get(d['id'])['pending_returns']
    assert store.issue(d['id'],issue['id'])['returns'][0]['acknowledged_at']


def test_return_and_ack_survive_restart_and_are_idempotent(store,svc):
    d=failed_check(store,svc);issue=store.report_issue(d['id'],d['version'],d['result_at'])
    first=store.return_to_warehouse(d['id'],issue['id'],'Доберите товар','idempotent1')
    with ThreadPoolExecutor(2) as pool:
        repeated=list(pool.map(lambda _:WarehouseStore(store.path).return_to_warehouse(d['id'],issue['id'],'Доберите товар','idempotent1'),range(2)))
    assert all(r['id']==first['id'] and r['duplicate'] for r in repeated)
    assert len(WarehouseStore(store.path).get(d['id'])['warehouse_replies'])==1
    with pytest.raises(WarehouseError):store.return_to_warehouse(d['id'],issue['id'],'Другой текст','idempotent1')
    with ThreadPoolExecutor(2) as pool:
        list(pool.map(lambda _:WarehouseStore(store.path).acknowledge_return(d['id'],first['id']),range(2)))
    assert not WarehouseStore(store.path).get(d['id'])['warehouse_replies']
    actions=[e['action'] for e in store.issue_events(d['id'],issue['id'])['events']]
    assert actions.count('returned_to_warehouse')==1 and actions.count('warehouse_return_acknowledged')==1


def test_unread_reply_does_not_block_scans_and_can_reply_to_closed_issue(store,svc):
    d=failed_check(store,svc);issue=store.report_issue(d['id'],d['version'],d['result_at'])
    store.resolve_issue(d['id'],issue['id'],'Проверено')
    store.return_to_warehouse(d['id'],issue['id'],'Можно собирать дальше','already_closed')
    add(store,d,code('INTRODUCEDQ100B'))
    assert len(store.get(d['id'])['scans'])==2
    assert store.get(d['id'])['pending_returns']==1
    assert asyncio.run(check_document(store,svc,d['id']))['can_complete']


def test_return_requires_comment_active_document_and_correct_issue(store,svc):
    d=failed_check(store,svc);issue=store.report_issue(d['id'],d['version'],d['result_at'])
    other=document(store)
    with pytest.raises(WarehouseError):store.return_to_warehouse(d['id'],issue['id'],' ','empty')
    with pytest.raises(WarehouseError):store.return_to_warehouse(other['id'],issue['id'],'Текст','wrong_doc')
    with pytest.raises(WarehouseError):store.acknowledge_return(d['id'],'missing')
    fresh=store.get(d['id']);store.delete(d['id'],fresh['version'],fresh['revision'])
    with pytest.raises(WarehouseError):store.return_to_warehouse(d['id'],issue['id'],'Текст','deleted')


def test_return_api_authentication_and_reply_identity(clients,store,svc):
    a,b=clients;h={'X-Admin-Key':'test-admin'}
    d=failed_check(store,svc);issue=store.report_issue(d['id'],d['version'],d['result_at'])
    endpoint='/api/admin/warehouse/documents/'+d['id']+'/issues/'+issue['id']+'/return'
    body={'resolution':'Замените ошибочный КМ','request_id':'test-request'}
    assert b.post(endpoint,json=body).status_code==403
    a.post('/api/admin/session',json={'key':'test-admin'})
    assert a.post(endpoint,json=body).status_code==403  # cookie alone is insufficient for a write
    reply=a.post(endpoint,json=body,headers={'X-Admin-Request':'1'}).json()
    public='/api/warehouse/documents/'+d['id']
    assert b.get(public).json()['warehouse_replies'][0]['id']==reply['id']
    other=document(store)
    assert b.post('/api/warehouse/documents/'+other['id']+'/returns/'+reply['id']+'/acknowledge').status_code==404
    assert b.post(public+'/returns/'+reply['id']+'/acknowledge').status_code==200
    assert not b.get(public).json()['pending_returns']


# Runtime application settings: persisted policies, concurrent admins and API enforcement.
def configure(store, **changes):
    previous = store.application_settings()
    return store.update_settings(previous['revision'], {**previous['values'], **changes})


def test_settings_persist_and_conflicting_admin_cannot_overwrite(store):
    old = store.application_settings()
    updated = configure(store, manual_input=False, receipt_default_org=ORG2)
    restarted = WarehouseStore(store.path)
    assert restarted.application_settings() == updated
    assert restarted.settings_history()[0]['changes']['manual_input'] == {'before': True, 'after': False}
    with pytest.raises(WarehouseError, match='другим администратором'):
        restarted.update_settings(old['revision'], old['values'])
    assert configure(store, manual_input=False)['revision'] == updated['revision']
    with pytest.raises(WarehouseError, match='справочнике'):
        configure(store, receipt_default_org='9999999999')
    assert restarted.application_settings() == updated


@pytest.mark.parametrize('patch', [
    {'shipment_statuses': []}, {'shipment_statuses': ['UNKNOWN']},
    {'shipment_statuses': ['INTRODUCED', 'INTRODUCED']},
    {'verdict_unknown': 'ok'}, {'check_timeout_seconds': 29},
    {'check_timeout_seconds': 541}, {'manual_input': 'false'}, {'invented': True},
    {'receipt_missing_codes_threshold': -1}, {'receipt_missing_codes_threshold': 10001},
    {'receipt_missing_codes_threshold': 1.5}, {'receipt_missing_codes_threshold': True},
])
def test_settings_api_rejects_invalid_values_without_mutation(clients, store, patch):
    a, _ = clients
    old = store.application_settings()
    r = a.put('/api/admin/warehouse/settings', headers={'X-Admin-Key': 'test-admin'},
              json={'revision': old['revision'], 'values': {**old['values'], **patch}})
    assert r.status_code == 422
    assert store.application_settings() == old
    assert store.settings_history() == []


@pytest.mark.parametrize('threshold', [0, 20, 10000])
@pytest.mark.asyncio
async def test_receipt_missing_display_threshold_preserves_check_and_full_result(store, svc, threshold):
    d = document(store, 'receipt')
    checked = await check_document(store, svc, d['id'])
    configure(store, receipt_missing_codes_threshold=threshold)
    current = WarehouseStore(store.path).get(d['id'])
    assert current['runtime_settings']['values']['receipt_missing_codes_threshold'] == threshold
    assert current['result'] == checked['result']
    assert len(current['result']['missing']) == 1
    assert current['check_current'] and not current['can_complete']


def test_settings_session_csrf_and_concurrent_save(clients, store):
    a, b = clients
    path = '/api/admin/warehouse/settings'
    old = store.application_settings()
    body = {'revision': old['revision'], 'values': {**old['values'], 'worker_delete_scans': False}}
    assert a.put(path, json=body).status_code == 403
    assert a.post('/api/admin/session', json={'key': 'test-admin'}).status_code == 200
    assert a.put(path, json=body).status_code == 403
    assert a.put(path, json=body, headers={'X-Admin-Request': '1'}).status_code == 200
    assert b.put(path, json=body, headers={'X-Admin-Key': 'test-admin'}).status_code == 409
    payload = a.get(path).json()
    assert payload['runtime']['revision'] == 1 and payload['history'][0]['revision'] == 1
    assert 'admin_key' not in payload['runtime']['values']


def test_admin_navigation_link_is_hidden_but_direct_admin_page_remains_available(clients):
    client, _ = clients
    for path in ('/', '/warehouse?kind=shipment', '/admin'):
        response = client.get(path)
        assert response.status_code == 200
        assert 'data-mode="admin"' not in response.text
    assert 'id="login"' in client.get('/admin').text


@pytest.mark.asyncio
async def test_runtime_verdict_applies_single_batch_but_never_masks_api_or_structure_errors(store, svc):
    svc.policy_getter = lambda: store.application_settings()['values']
    configure(store, verdict_emitted='error', verdict_applied='warning')
    assert (await svc.scan(code('EMITTED01'))).verdict == 'error'
    batch = await svc.scan_batch([code('EMITTED01'), code('APPLIED01'), 'bad'])
    assert [r.verdict for r in batch] == ['error', 'warning', 'error']
    assert batch[-1].error_category
    configure(store, verdict_emitted='ok')
    assert (await svc.scan(code('EMITTED01'))).verdict == 'ok'
    missing = await svc.scan(code('NOTFOUND01'))
    assert missing.error_category and missing.verdict == 'error'


@pytest.mark.asyncio
async def test_policy_change_invalidates_unfinished_shipment_preserves_receipt_and_completed(store, svc):
    shipped = document(store); add(store, shipped)
    completed = await check_document(store, svc, shipped['id']); store.complete(shipped['id'])
    receipt = document(store, 'receipt'); add(store, receipt)
    await check_document(store, svc, receipt['id'])
    d = document(store); add(store, d, code('APPLIED01'))
    failed = await check_document(store, svc, d['id'])
    assert not failed['result']['ok']
    # A check that started under old rules must not authorize completion after a save.
    check_id, snapshot = store.begin_check(d['id'])
    configure(store, shipment_statuses=['INTRODUCED', 'APPLIED'])
    old_result = {**failed['result'], 'ok': True}
    late = store.finish_check(d['id'], check_id, snapshot['version'], old_result)
    assert not late['check_current'] and not late['can_complete']
    with pytest.raises(WarehouseError):
        store.complete(d['id'])
    checked = await check_document(store, svc, d['id'])
    assert checked['check_current'] and checked['result']['policy']['settings_revision'] == 1
    # APPLIED fixture has a different quantity, so do not loosen quantity matching.
    assert not checked['result']['codes'][0]['errors']
    assert store.get(receipt['id'])['check_current']
    assert store.get(shipped['id'])['status'] == 'completed'
    assert store.get(shipped['id'])['check_current']
    assert len(checked['scans']) == 1


@pytest.mark.asyncio
async def test_feature_switches_enforced_server_side_without_blocking_scans(clients, store, svc):
    a, _ = clients
    d = document(store); add(store, d, code('INTRODUCEDQ50'))
    checked = await check_document(store, svc, d['id'])
    configure(store, report_comment_required=True)
    report = {'version': checked['version'], 'result_at': checked['result_at'], 'message': ''}
    path = '/api/warehouse/documents/' + d['id']
    assert a.post(path + '/issues', json=report).status_code == 422
    report['message'] = 'Недостача'
    issue = a.post(path + '/issues', json=report).json()
    configure(store, report_issues=False, return_to_warehouse=False, worker_delete_scans=False)
    assert a.post(path + '/issues', json=report).status_code == 409
    assert a.delete(path + '/scans/' + checked['scans'][0]['id']).status_code == 409
    admin_path = '/api/admin/warehouse/documents/' + d['id'] + '/issues/' + issue['id'] + '/return'
    assert a.post(admin_path, headers={'X-Admin-Key':'test-admin'}, json={'resolution':'Продолжайте', 'request_id':'settings-test'}).status_code == 409
    assert a.post(path + '/scans', json={'code':code('INTRODUCEDQ50B')}).status_code == 200
    assert len(store.get(d['id'])['scans']) == 2


@pytest.mark.asyncio
async def test_auto_pn_switch_and_default_owner_affect_only_new_documents(store, svc):
    old = document(store)
    configure(store, auto_pn=False, strict_owner_default=False)
    d = document(store, rows=[['', GTIN, 100]])
    assert not d['strict_owner'] and store.get(old['id'])['strict_owner']
    async def forbidden(*args, **kwargs):
        pytest.fail('Product lookup called when automatic PN is disabled')
    svc.client.product_info = forbidden
    add(store, d)
    checked = await check_document(store, svc, d['id'])
    assert checked['lines'][0]['pn'] == ''
    assert checked['can_complete']


@pytest.mark.asyncio
async def test_check_timeout_and_receipt_template_use_saved_settings(clients, store, svc, monkeypatch):
    configure(store, check_timeout_seconds=30, receipt_default_org=ORG2)
    original_timeout, deadlines = asyncio.timeout, []
    def record_timeout(delay):
        deadlines.append(delay)
        return original_timeout(delay)
    monkeypatch.setattr(asyncio, 'timeout', record_timeout)
    d = document(store); add(store, d)
    checked = await check_document(store, svc, d['id'])
    assert 30 in deadlines and checked['can_complete']
    a, _ = clients
    response = a.get('/api/admin/warehouse/templates/receipt', headers={'X-Admin-Key':'test-admin'})
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    metadata = {r[0]:r[1] for r in book['Документ'].iter_rows(min_row=2,values_only=True)}
    assert metadata['ИНН организации'] == ORG2
