"""Standard XLSX templates, strict import and lossless OOXML strings."""
from __future__ import annotations

import io
import re
import posixpath
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime

import openpyxl
import xlsxwriter

from datamatrix import gtin_checksum_valid, parse, GS
from organizations import registry
from warehouse_store import WarehouseError

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_ROWS = 50000
_OOXML_ESCAPE = re.compile(r'_x([0-9A-Fa-f]{4})_')


def decode_excel_string(value):
    # Single pass preserves literal _xNNNN_ escaped by OOXML as _x005F_xNNNN_.
    return _OOXML_ESCAPE.sub(lambda m: chr(int(m[1], 16)), value) if isinstance(value, str) else value


def read_excel_strings(content):
    """Read original OOXML text: openpyxl removes _x005F_ before callers can decode GS.

    Decode exactly once from XML, preserving a literal `_x001D_` in crypto data.
    Map sheet name -> cell coordinate -> decoded string; numeric cells stay numeric.
    """
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rel_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        strings = []
        if 'xl/sharedStrings.xml' in z.namelist():
            root = ET.fromstring(z.read('xl/sharedStrings.xml'))
            strings = [decode_excel_string(''.join(t.text or '' for t in si.findall('.//s:t', ns)))
                       for si in root.findall('s:si', ns)]
        relationships = {r.attrib['Id']: r.attrib['Target'] for r in
                         ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
        workbook = ET.fromstring(z.read('xl/workbook.xml'))
        result = {}
        for sheet in workbook.findall('s:sheets/s:sheet', ns):
            target = relationships[sheet.attrib[rel_ns]]
            path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
            cells = {}
            for _, cell in ET.iterparse(io.BytesIO(z.read(path)), events=('end',)):
                if cell.tag != '{' + ns['s'] + '}c':
                    continue
                typ, coord = cell.get('t'), cell.get('r')
                if typ == 's':
                    v = cell.find('s:v', ns)
                    if v is not None:
                        cells[coord] = strings[int(v.text)]
                elif typ == 'inlineStr':
                    cells[coord] = decode_excel_string(''.join(t.text or '' for t in cell.findall('.//s:t', ns)))
                cell.clear()
            result[sheet.attrib['name']] = cells
        return result


def canonical_full(raw):
    s = raw.strip(' \r\n')
    for prefix in (']d2', ']C1', ']e0', '^]'):
        if s.startswith(prefix):
            s = s[len(prefix):]
            break
    return s.replace('{GS}', GS).replace('^]', GS).replace('\r', '').replace('\n', '')


def scan_data(raw):
    if not isinstance(raw, str) or len(raw) > 4096:
        raise WarehouseError('Код должен быть строкой длиной до 4096 символов')
    p = parse(raw)
    full = canonical_full(raw)
    if not p.structure_valid or not full.startswith('01' + (p.gtin or '') + '21'):
        raise WarehouseError(p.structure_error or 'Некорректная структура DataMatrix')
    if not p.serial or len(p.serial) > 20:
        raise WarehouseError('Некорректная длина серийного номера; проверьте передачу GS сканером')
    # 91EE is the existing parser's explicit no-GS scanner fallback.
    if p.tail and not p.tail.startswith(GS):
        full = '01' + p.gtin + '21' + p.serial + GS + p.tail
    return dict(cis=p.normalized_cis, gtin=p.gtin, raw_code=raw, full_code=full, has_crypto=p.has_crypto)


def positive_integer(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('Количество должно быть положительным целым числом')
    text = str(value).strip()
    if not re.fullmatch(r'[0-9]+(?:\.0+)?', text):
        raise ValueError('Количество должно быть положительным целым числом')
    n = int(text.split('.')[0])
    if n <= 0 or n > 2**31 - 1:
        raise ValueError('Количество должно быть от 1 до 2147483647')
    return n


def text_cell(cell, label, required=True, strings=None):
    if cell.data_type == 'f':
        raise ValueError(f'{label}: формулы не допускаются')
    value = strings.get(cell.coordinate, cell.value) if strings is not None else cell.value
    if value is None and not required:
        return ''
    if not isinstance(value, str) or (required and not value.strip()):
        raise ValueError(f'{label}: заполните текстовую ячейку (формат «Текст»); ведущие нули важны')
    return value.strip() if label != 'Код маркировки' else value


def make_xlsx(sheets):
    stream = io.BytesIO()
    book = xlsxwriter.Workbook(stream, {'in_memory': True, 'strings_to_formulas': False, 'strings_to_urls': False})
    head = book.add_format({'bold': True, 'bg_color': '#DCEAF5', 'border': 1})
    text = book.add_format({'num_format': '@'})
    for name, rows in sheets.items():
        ws = book.add_worksheet(name)
        ws.freeze_panes(1, 0)
        ws.set_column(0, 0, 28, text)
        ws.set_column(1, 1, 24, text)
        ws.set_column(2, 2, 75, text)
        ws.set_column(3, 3, 22, text)
        for r, row in enumerate(rows):
            for col, value in enumerate(row):
                fmt = head if r == 0 else text
                if isinstance(value, str):
                    if len(value) > 32767:
                        raise WarehouseError('Слишком длинная строка для Excel')
                    ws.write_string(r, col, value, fmt)
                elif value is not None:
                    ws.write(r, col, value, fmt)
    book.close()
    return stream.getvalue()


def template(kind, default_org=''):
    if kind not in ('shipment', 'receipt'):
        raise WarehouseError('Неизвестный тип документа')
    header = [('Поле', 'Значение'), ('Номер документа', ''), ('Дата документа', date.today().isoformat()),
              ('ИНН организации', default_org)]
    if kind == 'shipment':
        header += [('Получатель', ''), ('ИНН получателя', '')]
    columns = ['PN', 'GTIN', 'Количество'] if kind == 'shipment' else ['PN', 'GTIN', 'Код маркировки', 'Количество в КМ']
    return make_xlsx({'Документ': header, 'Товары': [columns]})


def import_xlsx(content, kind, org_override=None):
    if kind not in ('shipment', 'receipt'):
        raise WarehouseError('Неизвестный тип документа')
    if len(content) > MAX_FILE_SIZE:
        raise WarehouseError('Размер файла превышает 10 МБ')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
                raise WarehouseError('Слишком большой распакованный Excel')
        original_strings = read_excel_strings(content)
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except WarehouseError:
        raise
    except Exception as exc:
        raise WarehouseError('Не удалось прочитать XLSX-файл') from exc
    try:
        if not {'Документ', 'Товары'} <= set(book.sheetnames):
            raise WarehouseError('Используйте шаблон: нужны листы «Документ» и «Товары»')
        meta = {}
        for row in book['Документ'].iter_rows(min_row=2, max_row=20, max_col=2):
            key = row[0].value
            if key is not None:
                try:
                    if key == 'Дата документа' and isinstance(row[1].value, (date, datetime)):
                        meta[key] = row[1].value.date().isoformat() if isinstance(row[1].value, datetime) else row[1].value.isoformat()
                    else:
                        meta[key] = text_cell(row[1], str(key), required=False, strings=original_strings.get('Документ', {}))
                except ValueError as exc:
                    raise WarehouseError(f'Документ, строка {row[0].row}: {exc}') from exc
        for key in ['Номер документа', 'Дата документа'] + (['Получатель', 'ИНН получателя'] if kind == 'shipment' else []):
            if not meta.get(key):
                raise WarehouseError(f'Документ: не заполнено поле «{key}»')
        try:
            date.fromisoformat(meta['Дата документа'])
        except ValueError as exc:
            raise WarehouseError('Дата документа: используйте ГГГГ-ММ-ДД') from exc
        inn = org_override or meta.get('ИНН организации', '')
        org = registry.get(inn)
        if not org:
            raise WarehouseError('ИНН нашей организации отсутствует в справочнике')
        if kind == 'shipment' and not re.fullmatch(r'(?:[0-9]{10}|[0-9]{12})', meta['ИНН получателя']):
            raise WarehouseError('ИНН получателя должен содержать 10 или 12 цифр')
        header = dict(kind=kind, number=meta['Номер документа'], date=meta['Дата документа'],
                      org_inn=inn, org_name=org.name, recipient_name=meta.get('Получатель', ''),
                      recipient_inn=meta.get('ИНН получателя', ''))
        ws = book['Товары']
        columns = ['PN', 'GTIN', 'Количество'] if kind == 'shipment' else ['PN', 'GTIN', 'Код маркировки', 'Количество в КМ']
        first = next(ws.iter_rows(min_row=1, max_row=1, max_col=len(columns)))
        if [c.value for c in first] != columns:
            raise WarehouseError('Товары, строка 1: ожидаются столбцы ' + ' / '.join(columns))
        lines, expected, seen = {}, [], set()
        for n, row in enumerate(ws.iter_rows(min_row=2, max_col=len(columns)), 2):
            if n > MAX_ROWS + 1:
                raise WarehouseError('В файле больше 50000 строк')
            if all(c.value is None for c in row):
                continue
            try:
                pn, gtin = text_cell(row[0], 'PN', False, original_strings.get('Товары', {})), text_cell(row[1], 'GTIN', strings=original_strings.get('Товары', {}))
                if not gtin_checksum_valid(gtin):
                    raise ValueError('GTIN должен содержать 14 цифр с верной контрольной цифрой')
                if row[-1].data_type == 'f':
                    raise ValueError('Количество: формулы не допускаются')
                qty = positive_integer(row[-1].value)
                if gtin in lines and lines[gtin]['pn'] and pn and lines[gtin]['pn'] != pn:
                    raise ValueError('У одного GTIN указаны разные PN')
                if kind == 'receipt':
                    code = scan_data(text_cell(row[2], 'Код маркировки', strings=original_strings.get('Товары', {})))
                    if code['gtin'] != gtin:
                        raise ValueError('GTIN не соответствует коду маркировки')
                    if code['cis'] in seen:
                        raise ValueError('Дубликат КМ в документе')
                    seen.add(code['cis'])
                    expected.append({**code, 'quantity': qty})
                if gtin not in lines:
                    lines[gtin] = dict(gtin=gtin, pn=pn, pn_source='excel' if pn else '', quantity=0)
                if pn and not lines[gtin]['pn']:
                    lines[gtin].update(pn=pn, pn_source='excel')
                lines[gtin]['quantity'] += qty
            except (ValueError, WarehouseError) as exc:
                raise WarehouseError(f'Товары, строка {n}: {exc}') from exc
        if not lines:
            raise WarehouseError('Товары: нет заполненных строк')
        return header, list(lines.values()), expected
    finally:
        book.close()


def export_shipment(doc):
    if doc['kind'] != 'shipment' or doc['status'] != 'completed':
        raise WarehouseError('Выгрузка доступна только после завершения отгрузки', 409)
    pns = {l['gtin']: l['pn'] for l in doc['lines']}
    if any(not v for v in pns.values()):
        raise WarehouseError('Заполните PN для всех позиций перед выгрузкой')
    if any(not s['has_crypto'] for s in doc['scans']):
        raise WarehouseError('КМ без криптографической части: требуется полное сканирование')
    # raw_code is deliberately used, never api_cis/normalized_cis.
    return make_xlsx({'Коды маркировки': [['PN', 'GTIN', 'Код маркировки']] +
                     [[pns[s['gtin']], s['gtin'], s['raw_code']] for s in doc['scans']]})
